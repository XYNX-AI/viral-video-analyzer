"""Excel exports for single-video and batch analysis results."""

from __future__ import annotations

import os
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


VIRAL_LABELS = {
    "hook": "Hook",
    "pain_point": "Pain Point",
    "demonstration": "Demonstration",
    "proof": "Proof",
    "cta": "CTA",
    "usage_scenario": "使用场景",
    "person": "人物",
    "product_appearance_time": "产品出现时间",
    "before_after": "Before / After",
    "video_duration": "视频时长",
    "visual_structure": "画面结构",
    "voiceover": "Voiceover",
    "caption": "Caption",
}


def _autosize(ws) -> None:
    for column_cells in ws.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(width, 80)


def _title(ws, values: list[str]) -> None:
    ws.append(values)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"


def _as_dict(result: Any) -> dict:
    if hasattr(result, "model_dump"):
        return result.model_dump()
    return result


def export_result_to_excel(result: Any, output_path: str) -> str:
    """Write one FullAnalysisResult to a readable multi-sheet workbook."""
    data = _as_dict(result)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "视频信息"
    _title(summary, ["字段", "值"])
    meta = data.get("video_meta", {})
    for key, value in [
        ("任务 ID", data.get("task_id", "")),
        ("创建时间", data.get("created_at", "")),
        ("标题", meta.get("title", "")),
        ("作者", meta.get("author", "")),
        ("点赞数", meta.get("likes", "")),
        ("评论数", meta.get("comments", "")),
        ("分享数", meta.get("shares", "")),
        ("时长", meta.get("duration", "")),
    ]:
        summary.append([key, value])

    viral_sheet = workbook.create_sheet("爆款分析")
    _title(viral_sheet, ["字段", "分析结果"])
    for key, label in VIRAL_LABELS.items():
        viral_sheet.append([label, data.get("viral_analysis", {}).get(key, "未明确")])

    transcript_sheet = workbook.create_sheet("口播文案")
    _title(transcript_sheet, ["开始时间", "结束时间", "文案"])
    for segment in data.get("transcript", []):
        transcript_sheet.append([segment.get("start", ""), segment.get("end", ""), segment.get("text", "")])

    segments_sheet = workbook.create_sheet("分段分析")
    _title(segments_sheet, ["开始时间", "结束时间", "分段文案", "画面描述", "画面提示词", "文案提示词", "复刻提示词"])
    frames_sheet = workbook.create_sheet("画面分析")
    _title(frames_sheet, ["分段开始", "分段结束", "时间点", "景别", "运镜", "构图", "转场", "画面文字", "画面描述", "情绪", "关键元素"])
    for segment in data.get("segments", []):
        prompts = segment.get("ai_prompts", {})
        descriptions = [f.get("visual_description", "") for f in segment.get("frames", [])]
        segments_sheet.append([
            segment.get("start_time", ""),
            segment.get("end_time", ""),
            segment.get("transcript", "") or "",
            "；".join(item for item in descriptions if item),
            prompts.get("visual", ""),
            prompts.get("copywriting", ""),
            prompts.get("recreation", ""),
        ])
        for frame in segment.get("frames", []):
            frames_sheet.append([
                segment.get("start_time", ""),
                segment.get("end_time", ""),
                frame.get("timestamp", ""),
                frame.get("shot_type", ""),
                frame.get("camera_movement", ""),
                frame.get("composition", ""),
                frame.get("transition", ""),
                frame.get("text_overlay", "") or "",
                frame.get("visual_description", ""),
                frame.get("mood", ""),
                "；".join(frame.get("key_elements", [])),
            ])

    for sheet in workbook.worksheets:
        _autosize(sheet)
    workbook.save(output_path)
    return output_path


def export_batch_summary(results: list[dict], output_path: str) -> str:
    """Write one summary row per URL, including failed items."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "批量汇总"
    columns = ["序号", "任务 ID", "链接", "状态", "标题", "作者", "时长", *VIRAL_LABELS.values(), "错误"]
    _title(sheet, columns)

    for index, item in enumerate(results, start=1):
        data = item.get("result", {})
        meta = data.get("video_meta", {})
        viral = data.get("viral_analysis", {})
        sheet.append([
            index,
            item.get("task_id", ""),
            item.get("url", ""),
            item.get("status", ""),
            meta.get("title", ""),
            meta.get("author", ""),
            meta.get("duration", ""),
            *[viral.get(key, "未明确") for key in VIRAL_LABELS],
            item.get("error", ""),
        ])

    _autosize(sheet)
    workbook.save(output_path)
    return output_path
